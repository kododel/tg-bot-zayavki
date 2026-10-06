"""Выгрузка заявок в Excel.

openpyxl — библиотека синхронная, поэтому сборку файла вызываем через
asyncio.to_thread: бот не должен замирать на пару секунд, пока пишется книга.
"""

from __future__ import annotations

import asyncio
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from bot.db import STATUS_LABELS, Application, all_for_export, get_stats

HEADERS = [
    "№",
    "Дата",
    "Имя",
    "Телефон",
    "Комментарий",
    "Статус",
    "Клиент в Telegram",
    "ID пользователя",
]

HEADER_FILL = PatternFill("solid", fgColor="2F5597")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Ширина колонок подбиралась вручную: «Комментарий» тянется,
# остальные — по содержимому.
COLUMN_WIDTHS = [6, 18, 22, 18, 52, 14, 20, 16]


def _style_header(sheet: Worksheet) -> None:
    for index, title in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
        sheet.column_dimensions[get_column_letter(index)].width = COLUMN_WIDTHS[index - 1]
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"


def _fill_applications(sheet: Worksheet, applications: list[Application]) -> None:
    for row_index, app in enumerate(applications, start=2):
        values = [
            app.id,
            app.created_at_local.strftime("%d.%m.%Y %H:%M"),
            app.full_name,
            app.phone,
            app.comment or "",
            STATUS_LABELS.get(app.status, app.status),
            app.client_link,
            app.user_id,
        ]
        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=(col_index == 5),
                horizontal="center" if col_index in (1, 6) else "left",
            )

    if applications:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(applications) + 1}"


def _build_summary_sheet(sheet: Worksheet, summary: dict) -> None:
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 12

    title = sheet.cell(row=1, column=1, value="Сводка по заявкам")
    title.font = Font(bold=True, size=14)
    sheet.merge_cells("A1:B1")

    rows = [
        ("Всего заявок", summary["total"]),
        ("Новые", summary["new"]),
        ("В работе", summary["in_progress"]),
        ("Выполнены", summary["done"]),
        ("Отклонены", summary["rejected"]),
        ("За сегодня", summary["today"]),
        ("За 7 дней", summary["week"]),
    ]

    for offset, (label, value) in enumerate(rows, start=3):
        label_cell = sheet.cell(row=offset, column=1, value=label)
        label_cell.border = BORDER
        value_cell = sheet.cell(row=offset, column=2, value=value)
        value_cell.border = BORDER
        value_cell.alignment = Alignment(horizontal="center")
        if label == "Всего заявок":
            label_cell.font = Font(bold=True)
            value_cell.font = Font(bold=True)


def _build_workbook(applications: list[Application], summary: dict) -> Workbook:
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Заявки"
    _style_header(sheet)
    _fill_applications(sheet, applications)

    _build_summary_sheet(workbook.create_sheet("Сводка"), summary)
    return workbook


async def export_applications(export_dir: Path) -> Path:
    """Собирает .xlsx со всеми заявками и возвращает путь к файлу."""
    applications = await all_for_export()
    summary = await get_stats()

    export_dir = Path(export_dir)
    export_dir.mkdir(parents=True, exist_ok=True)

    stamp = datetime.now().strftime("%Y-%m-%d_%H-%M")
    path = export_dir / f"zayavki_{stamp}.xlsx"

    workbook = await asyncio.to_thread(_build_workbook, applications, summary)
    await asyncio.to_thread(workbook.save, path)
    return path
